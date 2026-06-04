from selenium import webdriver
from openpyxl import load_workbook
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC

# Launch browser
driver = webdriver.Edge()
driver.maximize_window()

# Load Excel file
wb = load_workbook(
    r'C:\Users\Bala\PycharmProjects\PythonProject1\usernamepassword1.xlsx'
)
sheet = wb.active

# Counters
total = 0
passed = 0
failed = 0

# Read test data
for row_num in range(2, sheet.max_row + 1):

    username = sheet.cell(row=row_num, column=1).value
    password = sheet.cell(row=row_num, column=2).value
    expected = sheet.cell(row=row_num, column=3).value

    driver.get("https://practicetestautomation.com/practice-test-login/")

    # Enter username
    driver.find_element(By.NAME, "username").clear()
    driver.find_element(By.NAME, "username").send_keys(username)

    # Enter password
    driver.find_element(By.NAME, "password").clear()
    driver.find_element(By.NAME, "password").send_keys(password)

    # Click login button
    driver.find_element(By.ID, "submit").click()

    # Determine actual result
    try:
        WebDriverWait(driver, 3).until(
            EC.url_contains("logged-in-successfully")
        )
        actual = "pass"

    except:
        actual = "fail"

    # Compare actual vs expected
    if actual == expected.strip().lower():

        print(f"Username: {username}, Password: {password} --> Test Passed")

        passed += 1
        sheet.cell(row=row_num, column=4).value = "PASS"

    else:

        print(f"Username: {username}, Password: {password} --> Test Failed")

        failed += 1
        sheet.cell(row=row_num, column=4).value = "FAIL"

        # Capture screenshot for failed test cases
        driver.save_screenshot(
            rf"C:\Users\Bala\PycharmProjects\PythonProject1\Screenshots\TC_{row_num}.png"
        )

    total += 1

# Save results back to the Excel
wb.save(
    r'C:\Users\Bala\PycharmProjects\PythonProject1\usernamepassword1.xlsx'
)

# Final Summary
print("\n===== FINAL TEST SUMMARY =====")
print(f"Total  : {total}")
print(f"Passed : {passed}")
print(f"Failed : {failed}")

# Close resources


import os

path = os.path.abspath(f"Screenshots/TC_{row_num}.png")
driver.save_screenshot(path)

print(f"Screenshot saved at: {path}")

wb.close()
driver.quit()
