# from selenium import webdriver
# from selenium.webdriver.common.by import By
# driver = webdriver.Edge()
# driver.get("https://www.facebook.com")
# print(driver.title)
# driver.quit()
# driver.find_element(By.XPATH, "//input[@id = 'facebook']")
from numpy.ma.extras import row_stack
##PROJECT 1--->LOGIN TESTING
# from selenium import webdriver
# from selenium.webdriver.common.by import By
# from selenium.webdriver.support.ui import WebDriverWait
# from selenium.webdriver.support import expected_conditions as EC
# driver = webdriver.Edge()
# driver.get("https://practicetestautomation.com/practice-test-login/")
# driver.find_element(By.ID, "username").send_keys("student")
# driver.find_element(By.ID, "password").send_keys("Password123")
# driver.find_element(By.CLASS_NAME, "btn").click()
# #moves to dashboard page from login page
# logout_btn = driver.find_element(By.LINK_TEXT, "Log out")
#
# assert logout_btn.is_displayed()
#
# print("Successfully logged in")
#
# driver.quit()

# #PROJECT 2 -->>  INVALID LOGIN TESTING
# from selenium import webdriver
# from selenium.webdriver.common.by import By
# from selenium.webdriver.support.ui import WebDriverWait
# from selenium.webdriver.support import expected_conditions as EC
# driver = webdriver.Edge()
# driver.get("https://practicetestautomation.com/practice-test-login/")
# driver.find_element(By.NAME, "username").send_keys("student")
# driver.find_element(By.NAME, "password").send_keys("pass2")
# driver.find_element(By.CLASS_NAME, "btn").click()
# error_msg = driver.find_element(By.ID, "error").text
# msg = error_msg
# assert all(word in msg for word in ['password', 'invalid'])
# print("login not successful")

#project 3-----> data driven testing
from selenium import webdriver
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.common.by import By
from openpyxl import load_workbook
from selenium.webdriver.support import expected_conditions as EC

total = 0
passed = 0
failed = 0

driver = webdriver.Edge()
driver.maximize_window()
wb = load_workbook(r'C:\Users\Bala\PycharmProjects\PythonProject1\usernamepassword1.xlsx')
sheet = wb.active

for row in sheet.iter_rows(min_row=2, values_only=True):
    username, password, expected = row

    driver.get("https://practicetestautomation.com/practice-test-login/")

    #enter username
    driver.find_element(By.NAME, 'username').clear()
    driver.find_element(By.NAME, 'username').send_keys(username)

    #enter password
    driver.find_element(By.NAME, 'password').clear()
    driver.find_element(By.NAME, 'password').send_keys(password)

    #click the submit button
    driver.find_element(By.CLASS_NAME, 'btn').click()

    current_url = driver.current_url



    try:
        WebDriverWait(driver, 3).until(
            EC.url_contains("logged-in-successfully")
        )
        actual = "Pass"

    except:
        actual = "Fail"


    if actual.strip().lower() == expected.strip().lower():
        print(f"Username: {username}, Password: {password} --> Test Passed")
        passed += 1
    else:
        print(f"Username: {username}, Password: {password} --> Test failed")
        failed += 1


    total+=1

    print('\n=====Test summary=====')
    print(f"total:{total}", f"passed:{passed}", f"failed:{failed}")







